"""Generate the 10-slide Medical Sentiment executive PowerPoint."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(22, 35, 64)
DARK = RGBColor(39, 45, 55)
MUTED = RGBColor(95, 105, 121)


def add_title(slide, heading: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.42), Inches(12), Inches(0.75))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = heading
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.72), Inches(1.08), Inches(11.8), Inches(0.4))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = MUTED


def add_bullets(slide, items: list[str], x=0.8, y=1.55, w=5.6, h=4.9, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(10)


def add_chart(slide, filename: str, x=6.65, y=1.5, w=5.85, h=4.9) -> None:
    png = ROOT / "visuals" / filename
    if not png.exists():
        svg = png.with_suffix(".svg")
        if svg.exists():
            from cairosvg import svg2png
            svg2png(url=str(svg), write_to=str(png))
    if png.exists():
        slide.shapes.add_picture(str(png), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def main() -> None:
    metrics = json.loads((ROOT / "reports/benchmark_metrics.json").read_text())
    test = {row["model"]: row for row in metrics["models"] if row["split"] == "test"}
    baseline = test["TF-IDF + Logistic Regression"]
    advanced = test["TF-IDF + Calibrated Linear SVM"]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.2))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "Medical Sentiment Dataset"
    paragraph.font.size = Pt(40)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    box = slide.shapes.add_textbox(Inches(0.82), Inches(2.55), Inches(11), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "Responsible seven-class mental-health NLP classification"
    paragraph.font.size = Pt(22)
    paragraph.font.color.rgb = RGBColor(210, 222, 241)

    specs = [
        ("2. Business Problem", "Support content analytics and qualified human review without claiming diagnosis", ["High-volume text is difficult to summarize consistently.", "Dataset labels are broad categories, not clinical assessments.", "Privacy, uncertainty, and human oversight are core requirements."], "safety_workflow.png"),
        ("3. Dataset", "Kaggle: Sentiment Analysis for Mental Health", ["53,043 raw rows and 3 columns", "Seven imbalanced labels", "362 published null statements", "Aggregated from multiple public datasets"], "class_distribution.png"),
        ("4. Exploratory Analysis", "Quality controls before modeling", ["Remove null and empty text", "Deduplicate normalized statements before splitting", "Report conflicting labels", "Review long-text outliers rather than deleting blindly"], "text_length_distribution.png"),
        ("5. Modeling", "Transparent baseline versus calibrated advanced classifier", ["Baseline: word TF-IDF + Logistic Regression", "Advanced: word/character TF-IDF + calibrated Linear SVM", "Stratified 70/15/15 split", "Macro F1 selects the model"], "top_terms.png"),
        ("6. Results", "Deterministic overlapping synthetic benchmark", [f"Baseline test macro F1: {baseline['macro_f1']:.3f}", f"Advanced test macro F1: {advanced['macro_f1']:.3f}", f"Validation-selected model: {metrics['recommended_model']}", "Official Kaggle rerun required"], "model_comparison.png"),
        ("7. Key Insights", "Class-level errors matter more than a single accuracy number", ["Small classes need macro recall and F1", "Confusions reflect overlapping language", "Confidence does not equal clinical certainty", "Low-confidence cases should abstain"], "per_class_f1.png"),
        ("8. Recommendations", "Govern the workflow, not just the model", ["Never automate diagnosis or crisis decisions", "Use qualified human review", "Minimize and protect text data", "Test source leakage, bias, calibration, and drift"], None),
        ("9. Future Work", "Move from portfolio benchmark to validated research", ["Source-aware validation", "Transformer challenger", "Out-of-distribution detection", "Prospective human labeling and governance review"], None),
    ]
    for heading, subtitle, bullets, chart in specs:
        slide = presentation.slides.add_slide(blank)
        add_title(slide, heading, subtitle)
        if chart:
            add_bullets(slide, bullets)
            add_chart(slide, chart)
        else:
            add_bullets(slide, bullets, x=1.0, y=1.7, w=11.2, h=4.7, size=21)

    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.9))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "10. Conclusion"
    paragraph.font.size = Pt(34)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.35), Inches(11.2), Inches(2.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "Responsible mental-health NLP combines leakage-aware engineering, class-level evaluation, calibrated uncertainty, privacy controls, and qualified human oversight."
    paragraph.font.size = Pt(27)
    paragraph.font.color.rgb = RGBColor(220, 230, 245)
    paragraph.alignment = PP_ALIGN.CENTER

    output = Path(__file__).with_name("Medical_Sentiment_Executive_Presentation.pptx")
    presentation.save(output)
    print(f"Saved {output}")


if __name__ == "__main__":
    main()
